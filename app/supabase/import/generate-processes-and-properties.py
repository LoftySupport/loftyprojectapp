#!/usr/bin/env python3
"""
Turn Amber's processes-and-properties workbook into the seed migration.

    python3 generate-processes-and-properties.py            writes ../migrations/0079_…sql
    python3 generate-processes-and-properties.py --report   prints what it decided and why
    python3 generate-processes-and-properties.py --compact PATH
        writes the same rows as multi-row VALUES lists to PATH — a quarter of the size,
        for pasting into a console that will not take 255 KB. Same effect, same
        idempotence; the readable file in migrations/ stays the record.

WHY A SCRIPT, AND WHY ITS OUTPUT IS CHECKED IN
----------------------------------------------
The migration has to replay from nothing on a database with no Python and no workbook,
so the SQL is a file in the repo. But 285 property rows and 107 schedule lines typed by
hand would be 285 chances to mistype a key, so the file is generated, and the generator
is kept beside the workbook it read. Re-run it and diff: the diff is the change.

WHAT IT DECIDES, AND WHAT IT REFUSES TO
---------------------------------------
Everything here is a translation of a cell, never a guess at one:

  * "unknown (no data)" in the type column stays `unknown`. 87 rows. Most are plainly
    dates and it is not this script's place to say so — the format is set in Setup ›
    Properties, by a person, and until then the slot holds nothing.
  * A blank department stays a null owning team. 130 rows.
  * Nothing gets an expected duration. The workbook gives none for any process; the
    per-property SLA numbers are carried as given, on the property.
  * Nothing is marked a milestone, required, or restricted. The columns exist for a
    person to fill.

Five things ARE decided here, each listed under --report so they can be argued with:

  1. Which process a property's "Process block" names, where the words differ
     (BLOCK_TO_PROCESS). "FCR" is "Footings Construction Report (FCR)"; "Lodged for
     Planning Approval" is "Planning Approval". One block, "PWA & Invoice", spans two
     processes and is split by row.
  2. Which process a predecessor/successor name refers to (GANTT_ALIASES). Those cells
     use the older 57-step schedule's vocabulary. Where a name has no confident match it
     is left unmapped and reported, not forced.
  3. Obvious misspellings in names (TYPOS): Ordererd, Receieved, Construciton, depost,
     PRACTICLE, RAINWATWATER, CLOTHSLINE. Labels are renameable in the app; these are
     fixed on the way in and listed.
  4. Four "Issued" rows carry the SAME successor list — "3 Concept Plans; 5 Soil (Bore
     Logs); 4 Site Survey; 2 PWA Deposit" — under PWA, Contract Signing, 2nd Deposit and
     Variation. It is true of PWA and a copy-paste on the other three (a variation does
     not precede the concept plan). A successor list is credited to the first process
     that carries it; later identical copies are reported and skipped.
  5. A dependency that runs backwards through the lifecycle or the stage groups — a
     Stage 1 process waiting on a Stage 3 one, an Acquisition process waiting on a
     Pre-construction one — is refused and reported. The sheet's order is the sheet's
     own statement of what comes first.

The construction schedule's predecessor cells were mangled by Excel — "3, 4, 5" became a
date. DECODE below turns them back (the rule is documented at the function), and the
predecessor and successor columns are read as one edge set, each side checked against
the other.
"""
import collections
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SRC = HERE / "lofty-processes-and-properties-2026-09-01.xlsx"
OUT = HERE.parent / "migrations" / "0079_the_workbook_of_1_september.sql"

REPORT = "--report" in sys.argv
notes: list[str] = []


def note(s: str) -> None:
    notes.append(s)


# ------------------------------------------------------------------ vocabularies
STAGES = {
    "aquisition and development": "Acquisition & Development",
    "aquistions & development": "Acquisition & Development",
    "aquistion & development": "Acquisition & Development",
    "pre-constructions": "Pre-construction",
    "pre-construction": "Pre-construction",
    "construction": "Construction",
    "maintenance": "Maintenance",
    "completed": "Completed",
    "closed": "Closed",
    "cancelled": "Cancelled",
}

FORMATS = {
    "enum": "single select",
    "user": "person",
    "url": "link",
    "number": "number",
    "date": "date",
    "text": "text",
    "yes/no": "checkbox",
    "currency": "currency",
    "unknown (no data)": "unknown",
}

TEAMS = {
    "acquisition & development": "acquisition_development",
    "design": "design",
    "finance": "finance",
    "scheduling": "scheduling",
    "selections": "selections",
    "estimating": "estimating",
    "construction admin": "construction_admin",
    "construction (site supervisor)": "construction",
    "construction": "construction",
}

TYPOS = {
    "Ordererd": "Ordered",
    "Receieved": "Received",
    "Construciton": "Construction",
    "depost": "deposit",
    "PRACTICLE": "PRACTICAL",
    "RAINWATWATER": "RAINWATER",
    "CLOTHSLINE": "CLOTHESLINE",
}

# Property "Process block" -> process name in the Processes sheet, where they differ.
# A value of None means the block is a heading with no process behind it (the stage
# gates, the contract value); those properties belong to the stage alone.
BLOCK_TO_PROCESS = {
    "new project": "Project Creation",
    "new job": "Job Creation",
    "(1 day)": None,
    "added to trello": None,
    "contract value": None,
    "stage gate": None,
    "prelim eer": "Preliminary EER",
    "prelim footings": "Preliminary Footings",
    "lodged for planning approval": "Planning Approval",
    "contracts/send finalised quote": "Contract Signing",
    "2nd deposit - contract deposit": "2nd Deposit (contract deposit)",
    "fcr": "Footings Construction Report (FCR)",
    "construction price check (cpc) sheet requested": "Construction Price Check (CPC) Sheet",
    "brc/kbs": "Building Rules Consent (BRC)",
    "da": "Development Approval (DA)",
    "final site inspection, confirmation of services": "Final Site Inspection",
    "released to construction": "Construction Release",
}

# "PWA & Invoice" is one block over two processes; the row decides.
PWA_INVOICE_SPLIT = {
    "contract issued": "PWA",
    "contract signed": "PWA",
    "deposit 1 paid": "Invoice",
    "$ amount paid": "Invoice",
}

# The predecessor/successor cells speak the older 57-step schedule's vocabulary.
# Mapped only where the match is not in doubt; the rest is reported as unmapped.
GANTT_ALIASES = {
    "pwa issued": "PWA",
    "pwa deposit": "PWA",
    "concept plans": "Concept Plan",
    "site survey": "Site Survey",
    "soil (bore logs)": "Soil - Bore Logs",
    "planning drawings ordered": "Planning Drawings",
    "planning drawings completed": "Planning Drawings",
    "prelim fcr ordered": "Preliminary Footings",
    "prelim eer": "Preliminary EER",
    "prelim eer pricing": "Preliminary EER",
    "prelim eer review": "Preliminary EER",
    "civil plan ordered": "Civil Plan",
    "civil plan checked & received": "Civil Plan",
    "planning approval": "Planning Approval",
    "fencing/retaining checks": "Retaining, Fencing & BOB",
    "working drawings ordered": "Working Drawings",
    "working drawings completed & checked": "Working Drawings",
    "working drawings signed": "Working Drawings",
    "final quote check": "Final Quote Check",
    "selections commence schemes": "Selections",
    "contract deposit paid": "2nd Deposit (contract deposit)",
    "framing drawings and layouts": "Framing and Trusses",
    "final eer ordered": "Finalised EER",
    "final eer checked": "Finalised EER",
    "final fcr ordered": "Footings Construction Report (FCR)",
    "fcr check": "Footings Construction Report (FCR)",
    "1st site inspection": "1st Site Inspection",
    "1st site inspection actioned": "1st Site Inspection",
    "lightweight verandah engineering": "Attached lightweight Verandah Engineering",
    "brc/kbs lodged": "Building Rules Consent (BRC)",
    "production estimate (cpc)": "Construction Price Check (CPC) Sheet",
    "pegging plan received": "Pegging Plan/Encroachment Plan",
    "beam design": "Beam Design",
    "development approval": "Development Approval (DA)",
    "finance approval": "Finance",
    "2nd site inspection actioned": "Final Site Inspection",
    "sa water docs": "SA Water",
}


# --------------------------------------------------------------------- helpers
def fix_typos(s: str) -> str:
    out = s
    for bad, good in TYPOS.items():
        if bad in out:
            out = out.replace(bad, good)
            note(f"typo: {bad!r} -> {good!r} in {s!r}")
    return out


def clean(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def slug(s: str) -> str:
    s = s.lower()
    s = re.sub(r"^\s*\d+\s*-\s*", "", s)          # "1 - Footings" -> "Footings"
    s = re.sub(r"\b1st\b", "first", s)
    s = re.sub(r"\b2nd\b", "second", s)
    s = re.sub(r"\b3rd\b", "third", s)
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    if not s or s[0].isdigit():
        s = "p_" + s
    return s


def q(s) -> str:
    """SQL literal, or NULL."""
    if s is None or s == "":
        return "null"
    return "'" + str(s).replace("'", "''") + "'"


def stage_of(cell: str, where: str) -> str | None:
    key = clean(cell).lower()
    if not key:
        return None
    if key not in STAGES:
        raise SystemExit(f"{where}: unknown lifecycle stage {cell!r}")
    return STAGES[key]


def teams_of(cell: str, where: str) -> list[str]:
    out = []
    for part in [p.strip() for p in clean(cell).split(",") if p.strip()]:
        if part.lower() not in TEAMS:
            raise SystemExit(f"{where}: unknown department {part!r}")
        t = TEAMS[part.lower()]
        if t not in out:
            out.append(t)
    return out


def decode_ids(v) -> list[tuple[int, int]]:
    """
    A predecessor/successor cell as (id, lag_days) pairs.

    The schedule's lists were mangled on the way through Excel: a cell that read
    "3, 4, 5" was parsed as a date. Two shapes come back:
      * year 2026 (this year, filled in by Excel for a d/m pair): "6, 8" -> 2026-08-06,
        so the day and the month are the two ids.
      * any other year: "3, 4, 5" -> 2003-04-05, so year-2000, month and day are the three.
    "93+10" is id 93 with a lag of 10 days. Plain numbers are ids.
    """
    if v is None or v == "":
        return []
    if isinstance(v, (int, float)):
        return [(int(v), 0)]
    if isinstance(v, dt.datetime):
        if v.year == 2026:
            return [(v.day, 0), (v.month, 0)]
        return [(v.year - 2000, 0), (v.month, 0), (v.day, 0)]
    out = []
    for tok in str(v).split(","):
        tok = tok.strip()
        if not tok:
            continue
        m = re.fullmatch(r"(\d+)\s*\+\s*(\d+)", tok)
        if m:
            out.append((int(m.group(1)), int(m.group(2))))
        else:
            out.append((int(float(tok)), 0))
    return out


# -------------------------------------------------------------------- read it
wb = openpyxl.load_workbook(SRC, data_only=True)

# ---- lifecycle sheet: only checked, never seeded — 0076 owns the vocabulary
lc = [[clean(c) for c in row] for row in wb["Lifecycle Stages"].iter_rows(values_only=True)][1:]
lifecycle = [stage_of(r[1], f"Lifecycle Stages!{i+2}") for i, r in enumerate(lc) if r[1]]
EXPECTED_LIFECYCLE = ["Acquisition & Development", "Pre-construction", "Construction",
                      "Maintenance", "Completed", "Closed", "Cancelled"]
if lifecycle != EXPECTED_LIFECYCLE:
    raise SystemExit(f"the workbook's lifecycle {lifecycle} is not the one 0076 installs {EXPECTED_LIFECYCLE}")

# ---- processes sheet
Process = collections.namedtuple("Process", "key name stage scope position import_ref")
processes: dict[str, Process] = {}        # by name
process_order: list[str] = []
for i, row in enumerate(wb["Processes"].iter_rows(values_only=True), start=1):
    if i == 1:
        continue
    _, stage_cell, name_cell, type_cell, *_ = row
    name = fix_typos(clean(name_cell))
    if not name:
        continue
    stage = stage_of(stage_cell, f"Processes!{i}")
    scope = clean(type_cell).lower() or None

    # Two rows the sheet leaves short, filled from the Properties sheet where the same
    # process appears with its stage and level. Both are reported.
    if name == "Building Rules Consent (BRC)" and stage is None:
        stage, scope = "Pre-construction", "job"
        note("Processes!34 Building Rules Consent (BRC): stage and level blank; taken from the Properties sheet's BRC/KBS block (Pre-construction, Job)")
    if name == "3 - External Cladding" and stage == "Pre-construction":
        stage = "Construction"
        note("Processes!47 '3 - External Cladding' says Pre-Construction; its six siblings and its own tasks say Construction — read as Construction")
    if stage is None or scope not in ("project", "job"):
        raise SystemExit(f"Processes!{i} {name!r}: stage {stage!r} / level {scope!r} incomplete")

    if name in processes:
        raise SystemExit(f"Processes!{i}: {name!r} appears twice")
    processes[name] = Process(slug(name), name, stage, scope, 0, f"Processes!{i}")
    process_order.append(name)

# Positions: sheet order within a stage, except the numbered Construction processes,
# which the sheet lists 1,2,4,3,5,6,7 and whose own numbers are the order.
by_stage: dict[str, list[str]] = collections.defaultdict(list)
for n in process_order:
    by_stage[processes[n].stage].append(n)
for stage, names in by_stage.items():
    if stage == "Construction":
        names.sort(key=lambda n: int(re.match(r"(\d+)", n).group(1)))
    for pos, n in enumerate(names, start=1):
        processes[n] = processes[n]._replace(position=pos)

process_by_key = {p.key: p for p in processes.values()}
if len(process_by_key) != len(processes):
    raise SystemExit("two processes slugged to the same key")

# ---- properties sheet, top half (rows until the blank separator)
Prop = collections.namedtuple(
    "Prop", "key label stage stage_group scope format teams sla process_name position import_ref")
props: list[Prop] = []
prop_keys: set[str] = set()
stage_groups: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
edges: set[tuple[str, str]] = set()          # (process depends on, process)
unmapped: collections.Counter = collections.Counter()
position_in: collections.Counter = collections.Counter()
succ_lists: dict[str, str] = {}                # successor text -> first process carrying it
reported_copies: set[tuple[str, str]] = set()

ws = wb["Properties"]
rows = list(ws.iter_rows(values_only=True))
split_at = next(i for i, r in enumerate(rows) if i > 0 and all(clean(c) == "" for c in r[:9]))

for i in range(1, split_at):
    r = rows[i]
    ref = f"Properties!{i+1}"
    num, stage_cell, group_cell, block_cell, _full, label_cell, type_cell, level_cell, dept_cell, sla_cell, pred_cell, succ_cell = (list(r) + [None] * 12)[:12]
    label = fix_typos(clean(label_cell))
    if not label:
        continue
    stage = stage_of(stage_cell, ref)
    block = fix_typos(clean(block_cell))
    fmt_cell = clean(type_cell).lower()
    if fmt_cell not in FORMATS:
        raise SystemExit(f"{ref}: unknown type {type_cell!r}")
    fmt = FORMATS[fmt_cell]
    scope = clean(level_cell).lower()
    if scope not in ("project", "job"):
        raise SystemExit(f"{ref}: level {level_cell!r}")
    teams = teams_of(dept_cell, ref)
    sla = int(sla_cell) if isinstance(sla_cell, (int, float)) else None

    # Which process, if any.
    bl = block.lower()
    if bl == "pwa & invoice":
        pname = PWA_INVOICE_SPLIT.get(label.lower())
        if pname is None:
            raise SystemExit(f"{ref}: PWA & Invoice row {label!r} not in the split table")
    elif bl in BLOCK_TO_PROCESS:
        pname = BLOCK_TO_PROCESS[bl]
    else:
        pname = block
    if pname is not None and pname not in processes:
        raise SystemExit(f"{ref}: block {block!r} -> process {pname!r} is not in the Processes sheet")

    # The stage group ("Stage 1") is the process's, taken as the commonest word its
    # properties use. A group that merely repeats the lifecycle stage is no group.
    group = clean(group_cell)
    if group.lower() in STAGES:
        group = ""
    if pname and group:
        stage_groups[pname][group] += 1

    pkey = processes[pname].key if pname else None
    lslug = slug(label)
    key = lslug if (pkey is None or lslug.startswith(pkey)) else f"{pkey}_{lslug}"
    if key in prop_keys:
        note(f"{ref}: {label!r} under {block!r} repeats an earlier row — skipped as a duplicate")
        continue
    prop_keys.add(key)

    owner = pname or f"stage:{stage}"
    position_in[owner] += 1
    props.append(Prop(key, label, stage, group or None, scope, fmt, teams, sla, pname,
                      position_in[owner], ref))

    # Dependencies between processes, from this row's predecessor/successor cells.
    if pname:
        for cell, direction in ((pred_cell, "pred"), (succ_cell, "succ")):
            text = clean(cell)
            if direction == "succ" and text:
                first = succ_lists.setdefault(text, pname)
                if first != pname:
                    if (text, pname) not in reported_copies:
                        reported_copies.add((text, pname))
                        note(f"{ref}: successor list {text!r} under {pname!r} is the same text as {first!r}'s — read as a copy-paste and skipped")
                    continue
            for item in [x.strip() for x in text.split(";") if x.strip()]:
                name = re.sub(r"^\d+\s+", "", item).lower()
                other = GANTT_ALIASES.get(name)
                if other is None:
                    unmapped[item] += 1
                    continue
                if other == pname:
                    continue
                edges.add((other, pname) if direction == "pred" else (pname, other))

# A dependency cannot run backwards through the lifecycle or the stage groups.
LIFECYCLE_RANK = {s: i for i, s in enumerate(EXPECTED_LIFECYCLE)}
GROUP_RANK = {"Stage 1": 1, "Stage 2": 2, "Stage 3": 3, "Variation": 4}

def rank(pname: str) -> tuple[int, int | None]:
    p = processes[pname]
    group = stage_groups[pname].most_common(1)[0][0] if stage_groups.get(pname) else None
    return (LIFECYCLE_RANK[p.stage], GROUP_RANK.get(group) if group else None)

def runs_backwards(a: str, b: str) -> bool:
    """True when b (the dependent) sits earlier than a (what it waits on)."""
    ra, rb = rank(a), rank(b)
    if rb[0] != ra[0]:
        return rb[0] < ra[0]
    return ra[1] is not None and rb[1] is not None and rb[1] < ra[1]

for a, b in sorted(edges):
    if runs_backwards(a, b):
        note(f"dependency {b!r} waits on {a!r}, which sits later in the process — runs backwards and is dropped")
edges = {(a, b) for a, b in edges if not runs_backwards(a, b)}

# Cycle check on the process graph — the trigger would refuse one, and it is better to
# know here which edges close a loop.
def find_cycle_edges(edges: set[tuple[str, str]]) -> set[tuple[str, str]]:
    """Return edges that participate in a cycle (empty when the graph is a DAG)."""
    succ = collections.defaultdict(set)
    for a, b in edges:
        succ[a].add(b)
    def reaches(a, b, seen):
        if a == b:
            return True
        for n in succ[a]:
            if n not in seen:
                seen.add(n)
                if reaches(n, b, seen):
                    return True
        return False
    return {(a, b) for a, b in edges if reaches(b, a, set())}

cyclic = find_cycle_edges(edges)
for a, b in sorted(cyclic):
    note(f"dependency {a!r} -> {b!r} closes a loop with the reverse path and is dropped")
edges -= cyclic

# ---- properties sheet, bottom half: the construction schedule
Task = collections.namedtuple("Task", "ref process_name name team days parent_ref position")
tasks: list[Task] = []
declared_pred: set[tuple[int, int, int]] = set()   # (task, depends_on, lag)
declared_succ: set[tuple[int, int, int]] = set()
task_refs: set[int] = set()
current_parent: dict[str, int | None] = {}
task_pos: collections.Counter = collections.Counter()

for i in range(split_at + 2, len(rows)):
    r = rows[i]
    ref_cell, _con, process_cell, _task_name_dup, _full, name_cell, _fmt, _level, team_cell, days_cell, pred_cell, succ_cell = (list(r) + [None] * 12)[:12]
    if not isinstance(ref_cell, (int, float)):
        continue
    ref = int(ref_cell)
    pname = fix_typos(clean(process_cell))
    if pname not in processes:
        raise SystemExit(f"Properties!{i+1}: task process {pname!r} unknown")
    name = fix_typos(clean(name_cell))
    teams = teams_of(team_cell, f"Properties!{i+1}")
    if len(teams) > 1:
        raise SystemExit(f"Properties!{i+1}: a task names two teams")
    team = teams[0] if teams else None
    days = int(days_cell) if isinstance(days_cell, (int, float)) else None
    preds = decode_ids(pred_cell)
    succs = decode_ids(succ_cell)

    # A summary line — no team, no neighbours — is the parent of the lines under it.
    is_summary = team is None and not preds and not succs
    if is_summary:
        current_parent[pname] = ref
        parent = None
        note(f"schedule line {ref} {name!r} has no team and no neighbours — read as the parent of the lines beneath it")
    else:
        parent = current_parent.get(pname)

    task_pos[pname] += 1
    tasks.append(Task(ref, pname, name, team, days, parent, task_pos[pname]))
    task_refs.add(ref)
    for other, lag in preds:
        declared_pred.add((ref, other, lag))
    for other, lag in succs:
        declared_succ.add((other, ref, lag))

task_edges = declared_pred | declared_succ
one_sided = (declared_pred ^ declared_succ)
note(f"schedule: {len(declared_pred)} predecessor edges, {len(declared_succ)} successor edges, "
     f"{len(one_sided)} declared on one side only — the union of both sides is used")
bad = [(a, b) for a, b, _ in task_edges if a not in task_refs or b not in task_refs or a == b]
if bad:
    raise SystemExit(f"schedule edges naming unknown lines: {bad}")
# Two lags on one edge would be a contradiction; keep the larger and say so.
merged: dict[tuple[int, int], int] = {}
for a, b, lag in task_edges:
    if (a, b) in merged and merged[(a, b)] != lag:
        note(f"schedule edge {b} -> {a} carries two lags ({merged[(a,b)]}, {lag}); the larger is kept")
    merged[(a, b)] = max(lag, merged.get((a, b), 0))
by_process_task = {t.ref: t.process_name for t in tasks}
cross = [(a, b) for (a, b) in merged if by_process_task[a] != by_process_task[b]]
for a, b in cross:
    lost = f" — its {merged[(a, b)]}-day lag is not carried onto the process dependency" if merged[(a, b)] else ""
    note(f"schedule edge {b} -> {a} crosses from {by_process_task[b]!r} to {by_process_task[a]!r}; "
         f"kept as a process dependency instead of a task one{lost}")
    edges.add((by_process_task[b], by_process_task[a]))
merged = {k: v for k, v in merged.items() if k not in set(cross)}
cyc = find_cycle_edges({(str(b), str(a)) for (a, b) in merged})
if cyc:
    raise SystemExit(f"the schedule has a cycle: {sorted(cyc)}")
cyclic2 = find_cycle_edges(edges)
for a, b in sorted(cyclic2):
    note(f"process dependency {a!r} -> {b!r} (from a cross-process schedule edge) closes a loop and is dropped")
edges -= cyclic2


# ------------------------------------------------------------------ compact
if "--compact" in sys.argv:
    target = Path(sys.argv[sys.argv.index("--compact") + 1])
    c: list[str] = []
    cw = c.append
    cw("-- 0079 — the workbook of 1 September (compact form of the checked-in migration; same rows, same idempotence)")
    cw("create unique index if not exists process_tasks_one_per_import_ref on process_tasks (process_id, process_task_import_ref) where process_task_import_ref is not null;")
    cw("insert into processes (process_key, process_name, process_stage, process_stage_group, process_scope, process_position, process_import_ref) values")
    cw(",\n".join(f"({q(processes[n].key)},{q(processes[n].name)},{q(processes[n].stage)},{q(stage_groups[n].most_common(1)[0][0] if stage_groups.get(n) else None)},{q(processes[n].scope)},{processes[n].position},{q(processes[n].import_ref)})" for n in process_order) + " on conflict (process_key) do nothing;")
    cw("insert into process_dependencies (process_id, depends_on_process_id) select x.process_id, y.process_id from (values")
    cw(",\n".join(f"({q(processes[b].key)},{q(processes[a].key)})" for a, b in sorted(edges)) + ") e(k, dk) join processes x on x.process_key = e.k join processes y on y.process_key = e.dk on conflict do nothing;")
    cw("insert into property_defs (property_def_key, property_def_label, property_def_scope, property_def_stage, property_def_owning_team, property_def_format, property_def_position, property_def_sla_days, property_def_import_ref) values")
    cw(",\n".join(f"({q(p.key)},{q(p.label)},{q(p.scope)},{q(p.stage)},{q(p.teams[0] if p.teams else None)},{q(p.format)},{p.position},{p.sla if p.sla is not None else 'null'},{q(p.import_ref)})" for p in props) + " on conflict (property_def_key) do nothing;")
    cw("insert into property_access (property_def_key, team_id, property_access_can_create, property_access_can_read, property_access_can_update, property_access_can_delete) values")
    cw(",\n".join(f"({q(p.key)},{q(t)},true,true,true,false)" for p in props for t in p.teams) + " on conflict do nothing;")
    cw("insert into process_properties (process_id, property_def_key, process_property_position) select x.process_id, e.k, e.pos from (values")
    cw(",\n".join(f"({q(processes[p.process_name].key)},{q(p.key)},{p.position})" for p in props if p.process_name) + ") e(pk, k, pos) join processes x on x.process_key = e.pk on conflict do nothing;")
    cw("insert into process_tasks (process_id, process_task_name, process_task_owning_team, process_task_expected_days, process_task_position, process_task_import_ref) select x.process_id, e.n, e.t, e.d, e.pos, e.ref from (values")
    cw(",\n".join(f"({q(processes[t.process_name].key)},{q(t.name)},{q(t.team)},{t.days if t.days is not None else 'null'}::smallint,{t.position},{t.ref})" for t in tasks) + ") e(pk, n, t, d, pos, ref) join processes x on x.process_key = e.pk on conflict do nothing;")
    cw("update process_tasks c set parent_process_task_id = par.process_task_id from (values")
    cw(",\n".join(f"({q(processes[t.process_name].key)},{t.ref},{t.parent_ref})" for t in tasks if t.parent_ref is not None) + ") e(pk, ref, pref) join processes x on x.process_key = e.pk join process_tasks par on par.process_id = x.process_id and par.process_task_import_ref = e.pref where c.process_id = x.process_id and c.process_task_import_ref = e.ref and c.parent_process_task_id is null;")
    cw("insert into process_task_dependencies (process_task_id, depends_on_process_task_id, process_task_dependency_lag_days) select x.process_task_id, y.process_task_id, e.lag from (values")
    cw(",\n".join(f"({q(processes[by_process_task[a]].key)},{a},{b},{lag})" for (a, b), lag in sorted(merged.items())) + ") e(pk, a, b, lag) join processes pp on pp.process_key = e.pk join process_tasks x on x.process_id = pp.process_id and x.process_task_import_ref = e.a join process_tasks y on y.process_id = pp.process_id and y.process_task_import_ref = e.b on conflict do nothing;")
    cw("do $$ declare n integer; begin")
    for table, expected, where in (
        ("processes", len(processes), "process_import_ref like 'Processes!%'"),
        ("property_defs", len(props), "property_def_import_ref like 'Properties!%'"),
        ("process_dependencies", len(edges), "true"),
        ("process_tasks", len(tasks), "process_task_import_ref is not null"),
        ("process_task_dependencies", len(merged), "true"),
    ):
        cw(f"  select count(*) into n from {table} where {where}; if n < {expected} then raise exception '{table}: expected at least {expected} seeded rows, found %', n; end if;")
    cw(f"  raise notice 'ok  workbook seeded: {len(processes)} processes, {len(props)} properties, {len(edges)} dependencies, {len(tasks)} schedule lines'; end $$;")
    target.write_text("\n".join(c) + "\n")
    print(f"compact: {target} ({target.stat().st_size} bytes)")
    sys.exit(0)

# ------------------------------------------------------------------- report
if REPORT:
    print(f"processes: {len(processes)}")
    for s in EXPECTED_LIFECYCLE:
        ps = [processes[n] for n in by_stage.get(s, [])]
        if ps:
            print(f"  {s}: " + ", ".join(f"{p.name} [{p.scope}]" for p in sorted(ps, key=lambda p: p.position)))
    print(f"properties: {len(props)} ({sum(1 for p in props if p.format == 'unknown')} with format unknown, "
          f"{sum(1 for p in props if not p.teams)} with no team, {sum(1 for p in props if p.process_name is None)} with no process)")
    print(f"process dependencies: {len(edges)}")
    for a, b in sorted(edges):
        print(f"  {b}  <-  {a}")
    print(f"unmapped predecessor/successor names ({len(unmapped)}):")
    for name, n in unmapped.most_common():
        print(f"  {n:3d}× {name}")
    print(f"schedule tasks: {len(tasks)}, edges: {len(merged)}")
    print("notes:")
    for n in notes:
        print("  - " + n)
    sys.exit(0)

# ---------------------------------------------------------------------- write
out: list[str] = []
w = out.append
w("-- =============================================================================")
w("-- 0079 — the workbook of 1 September: processes, properties and the construction schedule")
w("-- =============================================================================")
w("-- GENERATED by app/supabase/import/generate-processes-and-properties.py from")
w(f"-- {SRC.name}. Edit the workbook or the script, re-run, and commit the diff;")
w("-- do not edit this file by hand.")
w("--")
w("-- Every row here is a cell in Amber's workbook, translated and not guessed at: the")
w("-- format `unknown` is the sheet's own \"unknown (no data)\"; a null team is a blank")
w("-- department; no process has an expected duration because the sheet gives none. The")
w("-- decisions the script did make — which process a block names, which process an older")
w("-- schedule name refers to, seven spelling fixes — are listed at the end of this file")
w("-- and in `python3 generate-processes-and-properties.py --report`.")
w("--")
w("-- Idempotent: every insert is ON CONFLICT DO NOTHING on its natural key, so a re-run")
w("-- against a database that already carries the workbook changes nothing, and a row")
w("-- somebody has since edited in the app keeps their edit.")
w("-- =============================================================================")
w("")
w("-- Template lines are addressed by their schedule number, so parents and dependencies")
w("-- can be wired without knowing the uuid they were given.")
w("create unique index if not exists process_tasks_one_per_import_ref")
w("  on process_tasks (process_id, process_task_import_ref) where process_task_import_ref is not null;")
w("")
w("-- ------------------------------------------------------------------ processes")
for n in process_order:
    p = processes[n]
    group = stage_groups[n].most_common(1)[0][0] if stage_groups.get(n) else None
    w("insert into processes (process_key, process_name, process_stage, process_stage_group, process_scope, process_position, process_import_ref)")
    w(f"values ({q(p.key)}, {q(p.name)}, {q(p.stage)}, {q(group)}, {q(p.scope)}, {p.position}, {q(p.import_ref)})")
    w("on conflict (process_key) do nothing;")
w("")
w("-- --------------------------------------------------------- process_dependencies")
for a, b in sorted(edges):
    w("insert into process_dependencies (process_id, depends_on_process_id)")
    w(f"select x.process_id, y.process_id from processes x, processes y where x.process_key = {q(processes[b].key)} and y.process_key = {q(processes[a].key)}")
    w("on conflict do nothing;")
w("")
w("-- ------------------------------------------------------------- property_defs")
for p in props:
    owner = p.teams[0] if p.teams else None
    w("insert into property_defs (property_def_key, property_def_label, property_def_scope, property_def_stage, property_def_owning_team,")
    w("                           property_def_format, property_def_position, property_def_sla_days, property_def_import_ref)")
    w(f"values ({q(p.key)}, {q(p.label)}, {q(p.scope)}, {q(p.stage)}, {q(owner)}, {q(p.format)}, {p.position}, "
      f"{p.sla if p.sla is not None else 'null'}, {q(p.import_ref)})")
    w("on conflict (property_def_key) do nothing;")
w("")
w("-- ------------------------------------------------------------ property_access")
w("-- Every department the sheet names may record and read the property. Nothing is")
w("-- restricted and nobody may delete: those are decisions for a person in the app.")
for p in props:
    for t in p.teams:
        w("insert into property_access (property_def_key, team_id, property_access_can_create, property_access_can_read, property_access_can_update, property_access_can_delete)")
        w(f"values ({q(p.key)}, {q(t)}, true, true, true, false)")
        w("on conflict do nothing;")
w("")
w("-- ---------------------------------------------------------- process_properties")
for p in props:
    if p.process_name:
        w("insert into process_properties (process_id, property_def_key, process_property_position)")
        w(f"select process_id, {q(p.key)}, {p.position} from processes where process_key = {q(processes[p.process_name].key)}")
        w("on conflict do nothing;")
w("")
w("-- ------------------------------------------------------------- process_tasks")
w("-- Parents first, so the child rows can find them.")
for t in sorted(tasks, key=lambda t: (t.parent_ref is not None, t.ref)):
    pk = processes[t.process_name].key
    parent_sql = "null" if t.parent_ref is None else (
        f"(select process_task_id from process_tasks pt join processes pp using (process_id) "
        f"where pp.process_key = {q(pk)} and pt.process_task_import_ref = {t.parent_ref})")
    w("insert into process_tasks (process_id, parent_process_task_id, process_task_name, process_task_owning_team, process_task_expected_days, process_task_position, process_task_import_ref)")
    w(f"select process_id, {parent_sql}, {q(t.name)}, {q(t.team)}, {t.days if t.days is not None else 'null'}, {t.position}, {t.ref}")
    w(f"from processes where process_key = {q(pk)}")
    w("on conflict do nothing;")
w("")
w("-- -------------------------------------------------- process_task_dependencies")
for (a, b), lag in sorted(merged.items()):
    pk = processes[by_process_task[a]].key
    w("insert into process_task_dependencies (process_task_id, depends_on_process_task_id, process_task_dependency_lag_days)")
    w(f"select x.process_task_id, y.process_task_id, {lag}")
    w(f"from processes pp join process_tasks x on x.process_id = pp.process_id and x.process_task_import_ref = {a}")
    w(f"                  join process_tasks y on y.process_id = pp.process_id and y.process_task_import_ref = {b}")
    w(f"where pp.process_key = {q(pk)}")
    w("on conflict do nothing;")
w("")
w("-- ---------------------------------------------------------------------- proof")
w("do $$")
w("declare n integer;")
w("begin")
for table, expected, where in (
    ("processes", len(processes), "process_import_ref like 'Processes!%'"),
    ("property_defs", len(props), "property_def_import_ref like 'Properties!%'"),
    ("process_dependencies", len(edges), "true"),
    ("process_tasks", len(tasks), "process_task_import_ref is not null"),
    ("process_task_dependencies", len(merged), "true"),
):
    w(f"  select count(*) into n from {table} where {where};")
    w(f"  if n < {expected} then raise exception '{table}: expected at least {expected} seeded rows, found %', n; end if;")
w(f"  raise notice 'ok  workbook seeded: {len(processes)} processes, {len(props)} properties, {len(edges)} dependencies, {len(tasks)} schedule lines';")
w("end $$;")
w("")
w("-- ------------------------------------------------------- what the script decided")
for n in notes:
    w("-- " + n)
w("-- unmapped predecessor/successor names (left as prose in the workbook, not wired):")
for name, count in unmapped.most_common():
    w(f"--   {count}x {name}")

OUT.write_text("\n".join(out) + "\n")
print(f"wrote {OUT.relative_to(HERE.parent.parent.parent)}: {len(processes)} processes, {len(props)} properties, "
      f"{len(edges)} process dependencies, {len(tasks)} schedule lines, {len(merged)} task dependencies")

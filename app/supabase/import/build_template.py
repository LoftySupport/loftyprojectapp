"""
Build the Phase B job-import template.

Regenerated rather than hand-edited: the dropdowns are the database's own lists — teams,
stages, statuses, councils — and a spreadsheet whose valid values drifted from the schema
would send somebody off to fix 200 rows that were right.

    python3 build_template.py

Lists are literals here rather than a live query on purpose: this has to build without a
database connection, and `verify/seeds.sh` is what catches them drifting.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

TEAMS = [
    ("acquisition_development", "Acquisition & Development"),
    ("sales_admin", "Sales Admin"),
    ("design", "Design"),
    ("pre_construction_admin", "Pre-Construction Admin"),
    ("scheduling", "Scheduling"),
    ("selections", "Selections"),
    ("estimating", "Estimating"),
    ("construction", "Construction"),
    ("construction_admin", "Construction Admin"),
    ("finance", "Finance"),
    ("maintenance", "Maintenance"),
    ("lofty_general", "Lofty General"),
]
# The five lifecycle phases, and only those. This list held the old nine until 0035
# retyped `jobs.job_stage` to text with a check on the five — so seven of the values
# this dropdown offered were ones the database would refuse on insert, and the example
# row below used one of them. Kept honest by verify/seeds.sh, which reads this file.
STAGES = [
    "Acquisition & Development", "Pre-construction", "Construction",
    "Handover & Maintenance", "Closed",
]
STATUSES = ["on_track", "at_risk", "behind_schedule", "on_hold", "completed", "cancelled", "archived"]
PROJECT_TYPES = ["residential", "commercial", "development"]
STATES = ["SA", "NSW", "VIC", "QLD", "WA", "NT", "TAS", "ACT"]
COUNCILS = """City of Adelaide|Adelaide Hills Council|Adelaide Plains Council|Alexandrina Council|The Barossa Council|Barunga West Council|Berri Barmera Council|City of Burnside|Campbelltown City Council|District Council of Ceduna|City of Charles Sturt|Clare and Gilbert Valleys Council|District Council of Cleve|District Council of Coober Pedy|Coorong District Council|Copper Coast Council|District Council of Elliston|The Flinders Ranges Council|District Council of Franklin Harbour|Town of Gawler|Regional Council of Goyder|City of Holdfast Bay|Kangaroo Island Council|District Council of Karoonda East Murray|District Council of Kimba|Kingston District Council|Light Regional Council|Lower Eyre Council|District Council of Loxton Waikerie|City of Marion|Mid Murray Council|City of Mitcham|Mount Barker District Council|City of Mount Gambier|District Council of Mount Remarkable|Rural City of Murray Bridge|Naracoorte Lucindale Council|Northern Areas Council|City of Norwood Payneham & St Peters|City of Onkaparinga|District Council of Orroroo Carrieton|District Council of Peterborough|City of Playford|City of Port Adelaide Enfield|Port Augusta City Council|City of Port Lincoln|Port Pirie Regional Council|City of Prospect|Renmark Paringa Council|District Council of Robe|Municipal Council of Roxby Downs|City of Salisbury|Southern Limestone Coast Council|Southern Mallee District Council|District Council of Streaky Bay|Tatiara District Council|City of Tea Tree Gully|District Council of Tumby Bay|City of Unley|City of Victor Harbor|Wakefield Regional Council|Town of Walkerville|Wattle Range Council|City of West Torrens|City of Whyalla|Wudinna District Council|District Council of Yankalilla|Yorke Peninsula Council""".split("|")

FONT = "Arial"
INK   = "1A1A1A"
HEAD  = "1F3A5F"          # heading bar
FILLIN = "FFF6D6"         # you fill this in
DERIVED = "EAEFF4"        # the app works this out
EXAMPLE = "F2F7F2"        # the example row
RULE  = "C9D2DB"

thin = Side(style="thin", color=RULE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# (header, width, kind, help)
#   kind: "fill" you type it · "pick" dropdown · "calc" a formula · "opt" optional
COLUMNS = [
    ("site_group", 22, "fill",
     "THE HARD ONE. Your name for the site these jobs share — anything consistent, e.g. "
     "\"Corner St GG\". Every row with the same value becomes ONE project. The old system "
     "has no project key, so this column is the only thing that groups them, and a job in "
     "the wrong group silently inherits the wrong council and the wrong site facts."),
    ("jobs_on_site", 12, "calc",
     "Counted from site_group. Sanity check only — if a site you know has four lots reads 3 "
     "or 5, the grouping is wrong somewhere."),
    ("lot_sequence", 12, "fill",
     "1, 2, 3… in LOT order within the site, not old-number order. This decides the job "
     "number: lot_sequence 2 becomes -02. Lofty's old numbers are scrambled against the "
     "lots, and lot-versus-sequence confusion is permanent once numbers are issued."),
    ("job_number_old", 15, "fill",
     "The old five-digit number, e.g. 12345. Kept forever and searchable — old paperwork, "
     "SharePoint folders and invoices carry it. Must be unique across the whole sheet."),
    ("lot_number", 12, "opt",
     "As it appears on the plan of division. Just the number — \"3\", not \"Lot 3\"."),
    ("street_number", 13, "opt",
     "Once titles have issued. Leave blank if the lot has no street number yet — but a row "
     "needs at least one of lot_number or street_number."),
    ("street", 24, "fill", "Street name only, e.g. Corner Street."),
    ("unit_level", 14, "opt", "Anything above the street line, e.g. Unit 3."),
    ("suburb", 18, "fill", "e.g. Golden Grove."),
    ("state", 9, "pick", "SA unless it genuinely is not."),
    ("postcode", 11, "fill", "Four digits, as text. 0800 is Darwin, so leading zeros matter."),
    ("council", 30, "pick",
     "Required for every SA address. Pick from the list — the database holds these 68 exactly "
     "and rejects anything else."),
    ("project_type", 15, "pick",
     "A property of the SITE, so every row in one site_group must say the same thing."),
    ("owning_team", 24, "pick", "Who is accountable for this job right now."),
    ("stage", 28, "pick", "Where the job sits today."),
    ("job_status", 15, "pick", "on_track unless you know otherwise."),
    ("notes", 34, "opt",
     "Anything the import should know — \"grouping not certain\", \"address changed in 2024\". "
     "Read by a person, not loaded."),
]

wb = Workbook()

# ─────────────────────────────────────────── Jobs
ws = wb.active
ws.title = "Jobs"

HEADER_ROW = 5
EXAMPLE_ROW = 6
FIRST_DATA_ROW = 7
LAST_ROW = 306          # 300 rows of room; Lofty has ~200 live jobs

legend = [
    ("Lofty job import — spine review", 14, True, HEAD),
    ("One row per JOB. Fill the amber columns; the grey one is a formula, leave it alone.", 10, False, None),
    # Interpolated, not typed. This line said "row 4" while the example sits in row 6
    # and the How-to sheet said 6 — two instructions disagreeing about which row to delete.
    (f"Row {EXAMPLE_ROW} is an example — delete it before sending this back.", 10, False, None),
]
for i, (text, size, bold, colour) in enumerate(legend, start=1):
    c = ws.cell(row=i, column=1, value=text)
    c.font = Font(name=FONT, size=size, bold=bold, color=colour or INK)

for idx, (name, width, kind, help_text) in enumerate(COLUMNS, start=1):
    letter = get_column_letter(idx)
    ws.column_dimensions[letter].width = width

    h = ws.cell(row=HEADER_ROW, column=idx, value=name)
    h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=HEAD)
    h.alignment = Alignment(vertical="center", wrap_text=True)
    h.border = BOX
    h.comment = Comment(f"{name}\n\n{help_text}", "Lofty import", width=340, height=170)

ws.row_dimensions[HEADER_ROW].height = 30

example = {
    "site_group": "Corner St GG", "lot_sequence": 1, "job_number_old": "12345",
    "lot_number": "1", "street_number": "", "street": "Corner Street", "unit_level": "",
    "suburb": "Golden Grove", "state": "SA", "postcode": "5125",
    "council": "City of Tea Tree Gully", "project_type": "residential",
    "owning_team": "Acquisition & Development", "stage": "Acquisition & Development",
    "job_status": "on_track", "notes": "example row — delete me",
}

for idx, (name, _w, kind, _h) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=EXAMPLE_ROW, column=idx)
    if kind != "calc":
        c.value = example.get(name, "")
    c.font = Font(name=FONT, size=10, italic=True, color="4A6B4A")
    c.fill = PatternFill("solid", fgColor=EXAMPLE)
    c.border = BOX

for row in range(EXAMPLE_ROW, LAST_ROW + 1):
    for idx, (name, _w, kind, _h) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=row, column=idx)
        c.border = BOX
        if row > EXAMPLE_ROW:
            c.font = Font(name=FONT, size=10, color=INK)
            c.fill = PatternFill("solid", fgColor=DERIVED if kind == "calc" else FILLIN)
        if name == "postcode":
            c.number_format = "@"        # text, so 5125 stays 5125 and 0800 stays 0800
        if name in ("lot_number", "street_number", "job_number_old"):
            c.number_format = "@"

# jobs_on_site: COUNTIF over the site_group column. Blank when the row is blank, so an
# empty sheet does not read "0 jobs on this site" 300 times.
site_col = get_column_letter(1)
for row in range(EXAMPLE_ROW, LAST_ROW + 1):
    ws.cell(row=row, column=2).value = (
        f'=IF({site_col}{row}="","",'
        f'COUNTIF(${site_col}${EXAMPLE_ROW}:${site_col}${LAST_ROW},{site_col}{row}))'
    )

ws.freeze_panes = f"C{FIRST_DATA_ROW - 1}"
ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}{LAST_ROW}"

# ─────────────────────────────────────────── Valid values
vv = wb.create_sheet("Valid values")
vv["A1"] = "The lists the database accepts. Do not edit — the dropdowns read these."
vv["A1"].font = Font(name=FONT, size=10, bold=True, color=HEAD)

lists = [
    ("state", STATES), ("project_type", PROJECT_TYPES), ("job_status", STATUSES),
    ("owning_team", [label for _slug, label in TEAMS]), ("stage", STAGES),
    ("council", COUNCILS),
]
positions = {}
for col, (name, values) in enumerate(lists, start=1):
    letter = get_column_letter(col)
    vv.column_dimensions[letter].width = max(14, min(34, max(len(v) for v in values) + 2))
    h = vv.cell(row=3, column=col, value=name)
    h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=HEAD)
    for i, v in enumerate(values, start=4):
        cell = vv.cell(row=i, column=col, value=v)
        cell.font = Font(name=FONT, size=10, color=INK)
    positions[name] = f"'Valid values'!${letter}$4:${letter}${3 + len(values)}"

# The team slug beside its name, because the import writes the slug and somebody will ask.
slug_col = len(lists) + 2
vv.column_dimensions[get_column_letter(slug_col)].width = 26
vv.column_dimensions[get_column_letter(slug_col + 1)].width = 26
for j, title in enumerate(("team name", "team_id (what is stored)")):
    h = vv.cell(row=3, column=slug_col + j, value=title)
    h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor=HEAD)
for i, (slug, label) in enumerate(TEAMS, start=4):
    vv.cell(row=i, column=slug_col, value=label).font = Font(name=FONT, size=10, color=INK)
    vv.cell(row=i, column=slug_col + 1, value=slug).font = Font(name=FONT, size=10, color="626B60")

# ─────────────────────────────────────────── dropdowns
for idx, (name, _w, kind, _h) in enumerate(COLUMNS, start=1):
    if kind != "pick":
        continue
    letter = get_column_letter(idx)
    dv = DataValidation(type="list", formula1=positions[name], allow_blank=True, showDropDown=False)
    dv.error = f"Pick a value from the list. The database rejects anything else in {name}."
    dv.errorTitle = "Not a valid value"
    dv.prompt = f"Choose from the {name} list on the Valid values sheet."
    ws.add_data_validation(dv)
    dv.add(f"{letter}{EXAMPLE_ROW}:{letter}{LAST_ROW}")

# ─────────────────────────────────────────── How to fill this in
how = wb.create_sheet("How to fill this in")
how.column_dimensions["A"].width = 3
how.column_dimensions["B"].width = 104

BLOCKS = [
    ("h", "How to fill this in"),
    ("p", "One row per job. About 200 live jobs; closed and cancelled ones can come later "
          "and need not be perfect."),
    ("h2", "The two things that are easy to get wrong"),
    ("n", "1.  site_group decides what a project IS.  The old system has no project key — its "
          "numbers are a flat sequence with nothing linking the jobs on one site, and they are "
          "not even contiguous. This column is the only thing that groups them. Every row "
          "sharing a site_group becomes one project with one project number."),
    ("i", "Why it matters more than it looks: project-level facts read through to every job on "
          "the project. A job in the wrong group silently shows the wrong council, the wrong "
          "developer and the wrong site facts — and nothing complains."),
    ("n", "2.  lot_sequence follows LOT order, not old-number order.  In Lofty's own example the "
          "old numbers are scrambled against the lots — Lot 3 is 12367, Lot 4 is 12356. Sorting "
          "by old number would make Lot 4 into job -03. Job numbers go on contracts, so that "
          "confusion is permanent."),
    ("h2", "If you are not sure two jobs belong together"),
    ("p", "Give them different site_groups. Each becomes a single-job project, which asserts "
          "nothing nobody verified. Project numbers are cheap; a wrong grouping is a lie that "
          "lives for years. Put why in the notes column."),
    ("h2", "What the app fills in, so you do not have to"),
    ("p", "•  project number — issued by the database, from 1000 up\n"
          "•  job number — 1042-01, composed from the project number and lot_sequence\n"
          "•  the consolidated address that search reads\n"
          "•  created/updated stamps"),
    ("h2", "Columns"),
    ("p", "Amber cells are yours. The grey jobs_on_site column is a formula — leave it alone; "
          "it counts rows sharing a site_group so a site you know has four lots showing 3 tells "
          "you something is miscoded.\n\n"
          "Every header has a comment with the detail — hover the little red corner."),
    ("h2", "Before you send it back"),
    ("p", f"•  Delete the example row (row {EXAMPLE_ROW})\n"
          "•  Check jobs_on_site against what you know for a handful of sites\n"
          "•  Every SA address needs a council\n"
          "•  Every row needs a lot_number or a street_number — either will do, not neither\n"
          "•  job_number_old must be unique across the sheet"),
    ("h2", "What happens then"),
    ("p", "The sheet loads verbatim into import_staging_jobs, and the projects and jobs are "
          "built from it. The raw rows stay in the database as the record of what was imported, "
          "so Phase C can read the same rows again for field values without anyone re-exporting "
          "anything."),
]

r = 2
for kind, text in BLOCKS:
    c = how.cell(row=r, column=2, value=text)
    if kind == "h":
        c.font = Font(name=FONT, size=15, bold=True, color=HEAD); r += 2; continue
    if kind == "h2":
        c.font = Font(name=FONT, size=11, bold=True, color=HEAD)
        r += 2; continue
    c.font = Font(name=FONT, size=10, italic=(kind == "i"),
                  color="4A5A6A" if kind == "i" else INK)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    how.row_dimensions[r].height = 14 * (1 + text.count("\n") + len(text) // 100)
    r += 2

wb.move_sheet("How to fill this in", offset=-2)
wb.save("lofty-job-import-template.xlsx")
print("wrote lofty-job-import-template.xlsx")
